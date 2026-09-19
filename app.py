import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
import io
import zipfile
import json

# ==========================================
# 1. إعدادات الصفحة و CSS للغة العربية (RTL)
# ==========================================
st.set_page_config(page_title="برنامج الرقمنة الإحترافي", page_icon="🎓", layout="wide")

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@400;700&display=swap');
    html, body, [class*="css"] {
        font-family: 'Tajawal', sans-serif;
        direction: rtl;
        text-align: right;
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 24px;
        direction: rtl;
    }
    .stButton>button {
        font-weight: bold;
    }
    .success-text { color: #27ae60; font-weight: bold; }
    .error-text { color: #e74c3c; font-weight: bold; }
    .warning-text { color: #f39c12; font-weight: bold; }
</style>
""", unsafe_allow_html=True)

# ==========================================
# 2. الثوابت وإعدادات الحالة (Session State)
# ==========================================

def get_default_settings():
    ar = [
        [0.0, 1.99, "ضعيف جدا"],
        [2.0, 3.99, "ضعيف يجب العمل أكثر"],
        [4.0, 4.99, "دون الوسط"],
        [5.0, 5.99, "فوق المتوسط"],
        [6.0, 7.5, "قريب من الجيد"],
        [7.51, 8.99, "جيد"],
        [9.0, 10.0, "ممتاز"],
    ]
    return {
        "العربية": [row[:] for row in ar],
        "الفرنسية": [
            [0.0, 1.99, "Très faible"], [2.0, 3.99, "Faible"],
            [4.0, 4.99, "Insuffisant"], [5.0, 5.99, "Passable"],
            [6.0, 7.5, "Assez bien"], [7.51, 8.99, "Bien"],
            [9.0, 10.0, "Excellent"],
        ],
        "الانجليزية": [
            [0.0, 1.99, "Very weak"], [2.0, 3.99, "Weak"],
            [4.0, 4.99, "Below average"], [5.0, 5.99, "Above average"],
            [6.0, 7.5, "Fairly good"], [7.51, 8.99, "Good"],
            [9.0, 10.0, "Excellent"],
        ],
        "الرياضة": [row[:] for row in ar],
    }

if 'obs_settings' not in st.session_state:
    st.session_state.obs_settings = get_default_settings()

if 'processed_files' not in st.session_state:
    st.session_state.processed_files = None

# ==========================================
# 3. الدوال المساعدة الأساسية
# ==========================================
def detect_subject_type(text):
    """التعرف على نوع المادة من خلال النص المستخرج"""
    t = str(text).lower()
    if any(k in t for k in ("فرنسية", "français", "fr")): return "الفرنسية"
    if any(k in t for k in ("نجليزي", "anglais", "ang")): return "الانجليزية"
    if any(k in t for k in ("بدنية", "رياضة", "sport", "eps")): return "الرياضة"
    return "العربية"

def find_header_and_data_rows(df):
    """دالة للبحث عن السطر الذي يحتوي على أسماء الأعمدة لتفادي أخطاء اختلاف الملفات"""
    for i in range(min(15, df.shape[0])):
        row_vals = [str(x).strip().lower() for x in df.iloc[i].values]
        if 'matricule' in row_vals or 'obs' in row_vals or 'nom' in row_vals:
            return i, i + 2 # السطر الإنجليزي (header)، و سطر بداية البيانات
    return 4, 5 # قيمة افتراضية إذا لم يجد شيئاً

def get_obs_col(df, header_row):
    obs_keywords = ["obs", "observation", "remarque", "الملاحظة", "ملاحظة", "ملاحظات", "الملاحظات"]
    for c in range(df.shape[1]):
        h = str(df.iloc[header_row, c]).strip().lower()
        if h in obs_keywords:
            return c
    return None

def get_mark_cols(df, header_row):
    excluded = {"", "nan", "nom", "date_n", "matricule", "prenom"}
    obs_keywords = {"obs", "observation", "remarque", "الملاحظة", "ملاحظة", "ملاحظات", "الملاحظات"}
    excluded.update(obs_keywords)
    
    mark_cols = []
    for c in range(df.shape[1]):
        h = str(df.iloc[header_row, c]).strip().lower()
        if h not in excluded:
            mark_cols.append(c)
    return mark_cols

def get_file_errors(excel_data_dict):
    report = {}
    total_err_count = 0
    for sname, df in excel_data_dict.items():
        h_row, d_row = find_header_and_data_rows(df)
        
        if df.shape[0] <= d_row:
            continue
            
        mark_cols = get_mark_cols(df, h_row)
        
        # لتسمية الأعمدة في التقرير، نستخدم السطر الذي يليه (الذي يحمل الأسماء بالعربية)
        lbl_row = h_row + 1 if h_row + 1 < df.shape[0] else h_row
        
        errs = []
        for r in range(d_row, df.shape[0]):
            for c in mark_cols:
                raw = df.iloc[r, c]
                cell_val = str(raw).strip() if pd.notna(raw) else ""
                col_lbl = str(df.iloc[lbl_row, c]).strip()
                
                if cell_val == "":
                    errs.append(f"السطر {r+1} | عمود '{col_lbl}': خانة فارغة")
                elif "," in cell_val:
                    errs.append(f"السطر {r+1} | عمود '{col_lbl}': فاصلة خاطئة → '{cell_val}'")
                else:
                    try:
                        v = float(cell_val)
                        if v < 0 or v > 20: 
                            errs.append(f"السطر {r+1} | عمود '{col_lbl}': قيمة خارج النطاق → {cell_val}")
                    except ValueError:
                        errs.append(f"السطر {r+1} | عمود '{col_lbl}': غير رقمي → '{cell_val}'")
        report[sname] = errs
        total_err_count += len(errs)
    return report, total_err_count

# ==========================================
# 4. دالة المعالجة والتصحيح (الأساسية)
# ==========================================
def process_workbook(file_buffer, insert_obs=False):
    xl = pd.ExcelFile(file_buffer)
    data_dict = {s: xl.parse(s, header=None) for s in xl.sheet_names}
    
    file_buffer.seek(0)
    wb = openpyxl.load_workbook(file_buffer)
    fill_clear = PatternFill(fill_type=None)
    
    corrections = 0
    for sname in wb.sheetnames:
        ws = wb[sname]
        df = data_dict.get(sname)
        
        h_row, d_row = find_header_and_data_rows(df)
        
        if df is None or df.shape[0] <= d_row:
            continue
            
        # استخراج اسم المادة من محتوى الملف (السطر الخامس) للحصول على دقة أعلى
        try:
            subj_txt = str(df.iloc[4, 0])
            subj_txt += " " + str(sname) # دمج اسم الشيت كإجراء احتياطي
        except:
            subj_txt = str(sname)
            
        det = detect_subject_type(subj_txt)
        rules = st.session_state.obs_settings.get(det, [])
        mcs = get_mark_cols(df, h_row)
        oc = get_obs_col(df, h_row)
        
        for r in range(d_row, df.shape[0]):
            xl_row = r + 1
            marks = []
            
            # 1. تصحيح الفواصل وجمع النقاط
            for c in mcs:
                cell = ws.cell(row=xl_row, column=c + 1)
                raw_val = cell.value
                if raw_val is not None:
                    val_str = str(raw_val).strip()
                    corrected = val_str.replace(",", ".")
                    
                    if corrected != val_str:
                        # تحويل النص إلى رقم حقيقي قبل حفظه في الإكسيل
                        try:
                            num_val = float(corrected)
                            cell.value = num_val
                        except:
                            cell.value = corrected
                            
                        cell.fill = fill_clear
                        corrections += 1
                        val_str = corrected
                    
                    try:
                        v = float(val_str)
                        if 0 <= v <= 20: 
                            marks.append(v)
                            cell.fill = fill_clear
                    except ValueError:
                        pass
                        
            # 2. إدراج الملاحظات (التقديرات)
            if insert_obs and (oc is not None) and len(marks) > 0:
                avg = round(sum(marks) / len(marks), 2)
                obs_text = ""
                
                for rule in rules:
                    try:
                        lo = float(rule[0])
                        hi = float(rule[1])
                        txt = str(rule[2])
                        if lo <= avg <= hi:
                            obs_text = txt
                            break
                    except:
                        continue
                
                # إذا لم يجد ملاحظة بسبب الثغرات، نعطيه أقرب ملاحظة
                if not obs_text and rules:
                    for rule in rules:
                        if avg < float(rule[0]):
                            obs_text = str(rule[2])
                            break
                    if not obs_text:
                        obs_text = str(rules[-1][2])

                if obs_text:
                    ws.cell(row=xl_row, column=oc + 1).value = obs_text
                    ws.cell(row=xl_row, column=oc + 1).fill = fill_clear
                    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, corrections

# ==========================================
# 5. واجهة المستخدم (التطبيقات والتبويبات)
# ==========================================
st.title("🎓 برنامج فحص وتصحيح نقاط الرقمنة الإحترافي")
st.markdown("**(يدعم معالجة الدفعات - نسخة الويب)**")

tab_main, tab_settings = st.tabs(["📂 فحص ومعالجة الملفات", "⚙️ إعدادات الملاحظات والتقديرات"])

# ----------------- التبويب الأول: المعالجة -----------------
with tab_main:
    uploaded_files = st.file_uploader(
        "اسحب ملفات الإكسيل (Excel) هنا، أو اضغط لاختيارها", 
        type=["xlsx", "xls"], 
        accept_multiple_files=True
    )

    if uploaded_files:
        st.info(f"تم رفع {len(uploaded_files)} ملف/ملفات.")
        
        col1, col2 = st.columns(2)
        with col1:
            if st.button("🚀 فحص الأخطاء فقط (تقرير)", use_container_width=True):
                st.session_state.processed_files = None
                with st.expander("📊 تقرير فحص الأخطاء الدقيق", expanded=True):
                    clean_files, error_files = 0, 0
                    for file in uploaded_files:
                        try:
                            xl = pd.ExcelFile(file)
                            data_dict = {s: xl.parse(s, header=None) for s in xl.sheet_names}
                            report, total_errs = get_file_errors(data_dict)
                            file.seek(0)
                            
                            st.markdown(f"**ملف:** `{file.name}`")
                            if total_errs == 0:
                                st.markdown("<span class='success-text'>✅ سليم تماماً وبدون أخطاء.</span>", unsafe_allow_html=True)
                                clean_files += 1
                            else:
                                st.markdown(f"<span class='error-text'>⚠️ يوجد {total_errs} خطأ:</span>", unsafe_allow_html=True)
                                for sheet, errs in report.items():
                                    if errs:
                                        st.write(f"- **اللسان ({sheet}):**")
                                        for e in errs:
                                            st.markdown(f"<span style='color:red; margin-right:20px;'>• {e}</span>", unsafe_allow_html=True)
                                error_files += 1
                            st.divider()
                        except Exception as e:
                            st.error(f"حدث خطأ أثناء قراءة {file.name}: {e}")
                    
                    st.success(f"نتيجة الفحص: {clean_files} ملفات سليمة | {error_files} ملفات بها أخطاء.")

        with col2:
            if st.button("✨ تصحيح آلي + إدراج الملاحظات", type="primary", use_container_width=True):
                with st.spinner('جاري معالجة الملفات وحقن الملاحظات...'):
                    processed_dict = {}
                    total_corrections = 0
                    for file in uploaded_files:
                        out_buffer, corrections = process_workbook(file, insert_obs=True)
                        processed_dict[file.name] = out_buffer
                        total_corrections += corrections
                        
                    st.session_state.processed_files = processed_dict
                    st.success(f"✅ تمت المعالجة بنجاح! تم إجراء {total_corrections} تصحيح وإدراج الملاحظات.")

        if st.session_state.processed_files:
            st.markdown("### 📥 تحميل الملفات المصححة")
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for fname, fbuffer in st.session_state.processed_files.items():
                    zip_file.writestr(f"{fname}", fbuffer.getvalue())
            
            st.download_button(
                label="📦 تحميل جميع الملفات (ZIP)",
                data=zip_buffer.getvalue(),
                file_name="الملفات_المصححة.zip",
                mime="application/zip",
                type="primary"
            )

# ----------------- التبويب الثاني: الإعدادات -----------------
with tab_settings:
    st.markdown("### ⚙️ إعدادات الملاحظات ومجالات النقاط")
    st.info("يمكنك تعديل المجالات والملاحظات مباشرة داخل الجداول أدناه. سيتم حفظ التعديلات في الجلسة الحالية.")
    
    settings = st.session_state.obs_settings
    subjects = list(settings.keys())
    
    edited_settings = {}
    cols = st.columns(2)
    
    for i, subj in enumerate(subjects):
        with cols[i % 2]:
            st.markdown(f"**{subj}**")
            df = pd.DataFrame(settings[subj], columns=["من (Min)", "إلى (Max)", "الملاحظة"])
            edited_df = st.data_editor(df, num_rows="dynamic", key=f"editor_{subj}", use_container_width=True)
            edited_settings[subj] = edited_df.values.tolist()
            
    if st.button("💾 حفظ التعديلات", type="primary"):
        st.session_state.obs_settings = edited_settings
        st.success("تم حفظ الإعدادات بنجاح!")
        
    st.divider()
    st.markdown("#### 📤/📥 تصدير واستيراد الإعدادات")
    scol1, scol2 = st.columns(2)
    with scol1:
        json_settings = json.dumps(st.session_state.obs_settings, ensure_ascii=False, indent=2)
        st.download_button(
            label="📤 تصدير الإعدادات (ملف TXT)",
            data=json_settings,
            file_name="اعدادات_الملاحظات.txt",
            mime="text/plain"
        )
    with scol2:
        settings_file = st.file_uploader("📥 استيراد الإعدادات (ملف TXT)", type=["txt"])
        if settings_file is not None:
            try:
                imported_settings = json.loads(settings_file.getvalue().decode('utf-8'))
                if st.button("تأكيد الاستيراد"):
                    st.session_state.obs_settings = imported_settings
                    st.success("تم الاستيراد بنجاح! سيتم تحديث الجداول.")
                    st.rerun()
            except Exception as e:
                st.error("صيغة الملف غير صحيحة.")

st.markdown("<hr><center><small style='color:grey;'>By: BkHouche | تم التحويل إلى نسخة الويب بواسطة Streamlit</small></center>", unsafe_allow_html=True)
